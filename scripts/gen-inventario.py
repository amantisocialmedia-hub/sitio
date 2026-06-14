# Genera AMANTI_Inventario_Maestro.xlsx desde data/products.json (con miniaturas embebidas)
import json, shutil
from pathlib import Path
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parent.parent
data = json.loads((ROOT / "data" / "products.json").read_text(encoding="utf-8"))
products = data["products"]
cats = {c["slug"]: c["name"] for c in data["categories"]}

WINE = "5E1B33"; CREAM = "FCF6F0"; YELLOW = "FFF2A8"
hdr_fill = PatternFill("solid", fgColor=WINE)
cream_fill = PatternFill("solid", fgColor=CREAM)
yellow_fill = PatternFill("solid", fgColor=YELLOW)
hdr_font = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="E5D8CC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

THUMB = 80
thumbs_dir = ROOT / "_thumbs"
if thumbs_dir.exists():
    shutil.rmtree(thumbs_dir)
thumbs_dir.mkdir()

def make_thumb(slug):
    """Crea miniatura 80x80 (proporción + fondo blanco) de la primera foto. Devuelve ruta o None."""
    src = ROOT / "public" / "products" / f"{slug}_1.jpg"
    if not src.exists():
        return None
    try:
        im = PILImage.open(src).convert("RGB")
    except Exception:
        return None
    im.thumbnail((THUMB, THUMB), PILImage.LANCZOS)
    canvas = PILImage.new("RGB", (THUMB, THUMB), "white")
    canvas.paste(im, ((THUMB - im.width) // 2, (THUMB - im.height) // 2))
    out = thumbs_dir / f"{slug}.jpg"
    canvas.save(out, "JPEG", quality=70, optimize=True)
    return out

wb = openpyxl.Workbook()

# ---------- Hoja Inventario ----------
ws = wb.active
ws.title = "Inventario"
cols = ["#", "Foto", "Categoría", "Producto", "Marca", "Precio", "Precio antes",
        "¿Tiene foto?", "# Fotos", "Slug", "Subtítulo", "Descripción"]
PRICE_COLS = (6, 7); DESC_COL = 12
ws.append(cols)
for ci, _ in enumerate(cols, 1):
    c = ws.cell(row=1, column=ci)
    c.fill = hdr_fill; c.font = hdr_font; c.border = border
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

rows = sorted(products, key=lambda p: (cats.get(p["category"], p["category"]).lower(), p["name"].lower()))
embedded = 0
for i, p in enumerate(rows, 1):
    imgs = p.get("images") or []
    has = len(imgs) > 0
    ws.append([
        i, None,
        cats.get(p["category"], p["category"]),
        p["name"], p.get("brand", ""),
        p.get("price", 0),
        p["compareAt"] if p.get("compareAt", 0) > 0 else None,
        "Sí" if has else "No",
        len(imgs), p["slug"], p.get("subtitle", ""), p.get("description", ""),
    ])
    r = ws.max_row
    fill = yellow_fill if not has else (cream_fill if i % 2 == 0 else None)
    for ci in range(1, len(cols) + 1):
        cell = ws.cell(row=r, column=ci)
        cell.border = border
        cell.alignment = Alignment(vertical="top", wrap_text=(ci == DESC_COL))
        if fill:
            cell.fill = fill
    for pc in PRICE_COLS:
        ws.cell(row=r, column=pc).number_format = '$#,##0'
    ws.row_dimensions[r].height = 62  # ~82px, encaja la miniatura de 80px
    if has:
        tp = make_thumb(p["slug"])
        if tp:
            img = XLImage(str(tp))
            ws.add_image(img, f"B{r}")
            embedded += 1

widths = [5, 13, 16, 30, 18, 11, 12, 11, 8, 30, 30, 70]
for ci, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(ci)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

# ---------- Hoja Resumen ----------
rs = wb.create_sheet("Resumen")
def hcell(cell, text):
    cell.value = text; cell.fill = hdr_fill; cell.font = hdr_font
    cell.alignment = Alignment(horizontal="center"); cell.border = border

con = sum(1 for p in products if (p.get("images") or []))
sin = len(products) - con
rs["A1"] = "AMANTI — Resumen de Inventario"; rs["A1"].font = Font(bold=True, size=14, color=WINE)
rs["A3"] = "Total de productos"; rs["B3"] = len(products)
rs["A4"] = "Con foto"; rs["B4"] = con
rs["A5"] = "Sin foto"; rs["B5"] = sin
for r in (3, 4, 5):
    rs[f"A{r}"].font = Font(bold=True)

hcell(rs["A7"], "Categoría"); hcell(rs["B7"], "Productos"); hcell(rs["C7"], "Precio promedio")
by = defaultdict(list)
for p in products:
    by[p["category"]].append(p)
row = 8
for slug, name in cats.items():
    items = by.get(slug, [])
    rs.cell(row=row, column=1, value=name)
    rs.cell(row=row, column=2, value=len(items))
    avg = round(sum(x["price"] for x in items) / len(items)) if items else 0
    c = rs.cell(row=row, column=3, value=avg); c.number_format = '$#,##0'
    for ci in (1, 2, 3):
        rs.cell(row=row, column=ci).border = border
    row += 1

caro = max(products, key=lambda p: p["price"])
barato = min(products, key=lambda p: p["price"])
rs.cell(row=row + 1, column=1, value="Más caro").font = Font(bold=True)
rs.cell(row=row + 1, column=2, value=f'{caro["name"]} — ${caro["price"]:,}')
rs.cell(row=row + 2, column=1, value="Más barato").font = Font(bold=True)
rs.cell(row=row + 2, column=2, value=f'{barato["name"]} — ${barato["price"]:,}')
rs.column_dimensions["A"].width = 22
rs.column_dimensions["B"].width = 30
rs.column_dimensions["C"].width = 16

out = ROOT / "AMANTI_Inventario_Maestro.xlsx"
wb.save(out)
shutil.rmtree(thumbs_dir)
mb = out.stat().st_size / 1024 / 1024
print("OK ->", out)
print(f"productos {len(products)} | con foto {con} | sin foto {sin} | miniaturas embebidas {embedded}")
print(f"peso archivo: {mb:.2f} MB")
